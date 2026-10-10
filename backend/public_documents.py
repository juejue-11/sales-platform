"""Read bounded public procurement attachments, keeping source facts independent of AI."""
from __future__ import annotations

import asyncio
import hashlib
import io
import ipaddress
import re
import socket
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import PurePosixPath
from urllib.parse import unquote, urljoin, urlsplit

import httpx

MAX_FILES = 4
MAX_BYTES = 12 * 1024 * 1024
MAX_PAGES = 80
MAX_TEXT = 240_000
MAX_ITEMS = 1000
_EXTENSIONS = (".pdf", ".doc", ".docx", ".xlsx", ".xls")
_COLUMN_NAMES = {
    "物料名称": "name", "材料名称": "name", "货物名称": "name", "商品名称": "name", "名称": "name",
    "规格型号": "specification", "规格": "specification", "型号": "specification",
    "材质": "material", "采购数量": "quantity", "数量": "quantity", "需求数量": "quantity",
    "计量单位": "unit", "单位": "unit", "交货期": "delivery_period", "交货时间": "delivery_period",
    "交付时间": "delivery_period", "交货地/交货港": "delivery_location", "交货地点": "delivery_location",
    "交付地点": "delivery_location", "物料编码": "code", "质量标准": "standard", "备注": "note",
}


class DocumentError(ValueError):
    pass


def _origin(url: str) -> tuple:
    parts = urlsplit(url)
    if parts.scheme not in ("https", "http") or not parts.hostname or parts.username or parts.password:
        raise DocumentError("invalid_source_url")
    try:
        port = parts.port or (443 if parts.scheme == "https" else 80)
    except ValueError:
        raise DocumentError("invalid_source_url") from None
    if port not in (80, 443):
        raise DocumentError("invalid_source_port")
    return parts.scheme, parts.hostname.lower(), port


async def validate_public_url(url: str, source_url: str) -> None:
    if _origin(url) != _origin(source_url):
        raise DocumentError("attachment_outside_source")
    host = urlsplit(url).hostname
    try:
        addresses = await asyncio.to_thread(socket.getaddrinfo, host, None, type=socket.SOCK_STREAM)
    except OSError:
        raise DocumentError("source_unreachable") from None
    if not addresses or any(not ipaddress.ip_address(address[4][0]).is_global for address in addresses):
        raise DocumentError("non_public_source")


class _Links(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links: list[dict] = []
        self.current: dict | None = None

    def handle_starttag(self, tag, attrs):
        if tag != "a":
            return
        attributes = dict(attrs)
        target = attributes.get("href", "")
        click = attributes.get("onclick", "")
        # Recognize this public-site link syntax; never execute JavaScript.
        match = re.search(r"ztbfjyz\(\s*['\"]([^'\"]+)['\"]", click)
        if match:
            target = match.group(1)
        if target and unquote(urlsplit(target).path).lower().endswith(_EXTENSIONS):
            self.current = {"url": target, "name": attributes.get("title", "")}
            self.links.append(self.current)

    def handle_data(self, data):
        if self.current and data.strip() not in self.current["name"]:
            self.current["name"] += data.strip()

    def handle_endtag(self, tag):
        if tag == "a":
            self.current = None


def discover_attachments(html: str, source_url: str) -> list[dict]:
    parser = _Links()
    parser.feed(html)
    found, seen = [], set()
    for link in parser.links:
        url = urljoin(source_url, link["url"])
        try:
            if _origin(url) != _origin(source_url):
                continue
        except DocumentError:
            continue
        if url not in seen:
            seen.add(url)
            found.append({"url": url, "name": link["name"] or unquote(PurePosixPath(urlsplit(url).path).name)})
    return found


def _cell(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _pdf(raw: bytes, document: dict) -> None:
    import pdfplumber

    pages, texts, items = [], [], []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        document["pages"] = len(pdf.pages)
        if len(pdf.pages) > MAX_PAGES:
            raise DocumentError("page_limit")
        headers = None
        for number, page in enumerate(pdf.pages, 1):
            body = page.extract_text() or ""
            row_texts = []
            for table in page.extract_tables():
                for raw_row in table:
                    row = [_cell(value) for value in raw_row]
                    meanings = {_COLUMN_NAMES.get(cell) for cell in row}
                    if "name" in meanings and "quantity" in meanings:
                        headers = row
                        continue
                    if not headers or len(row) != len(headers):
                        continue
                    values = dict(zip(headers, row))
                    facts = {_COLUMN_NAMES[key]: value for key, value in values.items() if key in _COLUMN_NAMES}
                    sequence = values.get("序") or values.get("序号")
                    if not facts.get("name") or (sequence and not sequence.isdigit()):
                        continue
                    quote = "；".join(f"{key}：{value}" for key, value in values.items() if value)
                    row_texts.append(quote)
                    items.append({
                        **{key: facts.get(key, "") for key in ("name", "specification", "material", "quantity", "unit",
                                                               "delivery_period", "delivery_location", "code", "standard", "note")},
                        "source_document_id": document["id"],
                        "record_id": document["record_id"], "page": number, "quote": quote,
                    })
                    if len(items) > MAX_ITEMS:
                        raise DocumentError("item_limit")
            text = body + ("\n表格逐行摘录：\n" + "\n".join(row_texts) if row_texts else "")
            pages.append({"page": number, "text": text})
            texts.append(f"第{number}页\n{text}")
            if sum(map(len, texts)) > MAX_TEXT:
                raise DocumentError("text_limit")
    document.update(text="\n\n".join(texts), page_texts=pages, items=items)
    document["status"] = "parsed" if document["text"].strip() else "unreadable"
    if document["status"] == "unreadable":
        document["error"] = "PDF无可提取文字，需要OCR或人工补充"


def parse_document(raw: bytes, *, record_id: int, url: str, name: str = "", mime_type: str = "") -> dict:
    digest = hashlib.sha256(raw).hexdigest()
    identifier = "doc-" + hashlib.sha256(f"{record_id}:{url}:{digest}".encode()).hexdigest()[:24]
    document = {"id": identifier, "record_id": record_id, "url": url,
                "name": name or unquote(PurePosixPath(urlsplit(url).path).name), "status": "unsupported",
                "pages": 0, "sha256": digest, "fetched_at": datetime.now(timezone.utc).isoformat(),
                "mime_type": mime_type, "byte_size": len(raw), "text": "", "page_texts": [], "items": []}
    if len(raw) > MAX_BYTES:
        document.update(status="failed", error="attachment_too_large")
        return document
    try:
        if raw.startswith(b"%PDF"):
            document["mime_type"] = "application/pdf"
            _pdf(raw, document)
        elif raw.startswith(b"PK") and urlsplit(url).path.lower().endswith(".docx"):
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(member.file_size for member in archive.infolist()) > 25 * 1024 * 1024:
                    raise DocumentError("expanded_document_limit")
            doc = Document(io.BytesIO(raw))
            lines = [p.text for p in doc.paragraphs]
            lines += ["；".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
            text = "\n".join(lines)
            if len(text) > MAX_TEXT:
                raise DocumentError("text_limit")
            document.update(text=text, status="parsed" if text.strip() else "unreadable")
        elif raw.lstrip().lower().startswith((b"<!doctype html", b"<html")):
            document.update(status="auth_required", error="附件返回网页，可能需要登录或验证")
        else:
            document["error"] = "该附件格式暂不能自动解析，需人工核对原文件"
    except DocumentError as exc:
        document.update(status="failed", error=str(exc), text="", page_texts=[], items=[])
    except Exception:
        document.update(status="unreadable", error="附件无法解析，请核对原文件", text="", page_texts=[], items=[])
    return document


async def _download(client: httpx.AsyncClient, url: str, source_url: str) -> tuple[bytes, str]:
    for _ in range(4):
        await validate_public_url(url, source_url)
        async with client.stream("GET", url, headers={"User-Agent": "ProcurementFactReader/1.0", "Accept": "*/*"}) as response:
            if response.status_code in (301, 302, 303, 307, 308):
                url = urljoin(url, response.headers.get("location", ""))
                continue
            if response.status_code in (401, 403):
                raise DocumentError("auth_required")
            if response.status_code != 200:
                raise DocumentError("download_failed")
            if int(response.headers.get("content-length", "0") or 0) > MAX_BYTES:
                raise DocumentError("attachment_too_large")
            chunks, size = [], 0
            async for chunk in response.aiter_bytes():
                size += len(chunk)
                if size > MAX_BYTES:
                    raise DocumentError("attachment_too_large")
                chunks.append(chunk)
            return b"".join(chunks), response.headers.get("content-type", "").split(";")[0].lower()
    raise DocumentError("redirect_limit")


async def fetch_record_documents(record: dict) -> list[dict]:
    source_url = record.get("source_url") or record.get("original_url") or record.get("canonical_url") or ""
    record_id = int(record.get("id") or record.get("record_id"))
    _origin(source_url)
    async with asyncio.timeout(45), httpx.AsyncClient(timeout=15, follow_redirects=False, trust_env=False) as client:
        raw, mime = await _download(client, source_url, source_url)
        if raw.startswith(b"%PDF"):
            return [await asyncio.to_thread(parse_document, raw, record_id=record_id, url=source_url, mime_type=mime)]
        # Public notices generally declare their charset; default to UTF-8 without speculative URL searching.
        encoding = "gb18030" if re.search(br"charset\s*=\s*['\"]?(?:gb2312|gbk|gb18030)", raw[:4096], re.I) else "utf-8"
        links = discover_attachments(raw.decode(encoding, errors="replace"), source_url)
        if len(links) > MAX_FILES:
            # Keep the omission visible instead of claiming all attachments were checked.
            omitted = {"url": source_url, "name": "其他附件", "error": "attachment_count_limit"}
        else:
            omitted = None
        documents = []
        for link in links[:MAX_FILES]:
            try:
                content, content_type = await _download(client, link["url"], source_url)
                document = await asyncio.to_thread(parse_document, content, record_id=record_id,
                                                  url=link["url"], name=link["name"], mime_type=content_type)
            except (DocumentError, httpx.HTTPError) as exc:
                document = parse_document(b"", record_id=record_id, url=link["url"], name=link["name"])
                code = str(exc) if isinstance(exc, DocumentError) else "source_unreachable"
                document.update(status="auth_required" if code == "auth_required" else "failed", error=code)
            documents.append(document)
        if omitted:
            document = parse_document(b"", record_id=record_id, url=omitted["url"], name=omitted["name"])
            document.update(status="failed", error=omitted["error"])
            documents.append(document)
        return documents
