"""Capture public attachments separately from the original notice body."""
import hashlib
import io
import ipaddress
import json
import re
import socket
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlsplit, unquote

import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader

from .legacy_word import WordConversionError, convert_legacy_word

MAX_BYTES = 12 * 1024 * 1024


def parse_content(raw, url, *, timeout=90):
    """Extract downloaded content without changing its original bytes/hash."""
    if len(raw) > MAX_BYTES:
        raise ValueError('attachment_too_large')
    extension = Path(unquote(urlsplit(url).path)).suffix.lower()
    if raw.lstrip().lower().startswith((b'<html', b'<!doctype html')):
        return dict(status='auth_required', error='attachment_returned_html', text='', pages=0)
    if not raw.startswith(b'%PDF') and extension == '.doc':
        raw = convert_legacy_word(raw, timeout=min(90, max(0.01, timeout)))
    if raw.startswith(b'%PDF'):
        reader = PdfReader(io.BytesIO(raw))
        pages = len(reader.pages)
        if pages > 80:
            raise ValueError('page_limit')
        texts, length = [], 0
        for page in reader.pages:
            text = page.extract_text() or ''
            length += len(text) + 1
            if length > 240000:
                raise ValueError('text_limit')
            texts.append(text)
        text = '\n'.join(texts)
        result = dict(text=text, pages=pages, status='parsed' if text.strip() else 'unreadable')
        if not text.strip():
            result['error'] = 'needs_ocr'
        return result
    if raw.startswith(b'PK') and extension == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(member.file_size for member in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError('expanded_document_limit')
        doc = Document(io.BytesIO(raw))
        text = '\n'.join([p.text for p in doc.paragraphs] +
                         ['；'.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
        if len(text) > 240000:
            raise ValueError('text_limit')
        return dict(text=text, pages=0, status='parsed' if text.strip() else 'unreadable')
    return dict(status='unsupported', error='format_not_supported', text='', pages=0)


def discover(response):
    if 'html' not in response.headers.get('Content-Type', ''):
        return []
    links, seen = [], set()
    origin = urlsplit(response.url)
    for anchor in BeautifulSoup(response.text, 'html.parser').find_all('a'):
        href = anchor.get('href', '')
        match = re.search(r"ztbfjyz\(\s*['\"]([^'\"]+)['\"]", anchor.get('onclick', ''))
        if match:
            href = match.group(1)
        url = urljoin(response.url, href)
        target = urlsplit(url)
        if (target.scheme, target.netloc) != (origin.scheme, origin.netloc):
            continue
        if not unquote(target.path).lower().endswith(('.pdf', '.doc', '.docx', '.xlsx', '.xls')) or url in seen:
            continue
        seen.add(url)
        links.append({'url': url, 'name': anchor.get_text('', strip=True) or unquote(Path(target.path).name)})
    return links


def _download(url, source_url, deadline=None):
    origin = urlsplit(source_url)
    deadline = deadline or time.monotonic() + 45
    for _ in range(4):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise ValueError('attachment_timeout')
        target = urlsplit(url)
        if (target.scheme, target.netloc) != (origin.scheme, origin.netloc) or target.username or target.password:
            raise ValueError('attachment_outside_source')
        if target.port not in (None, 80, 443):
            raise ValueError('invalid_port')
        addresses = socket.getaddrinfo(target.hostname, None, type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
            raise ValueError('non_public_source')
        with requests.get(url, timeout=(min(5,remaining), min(15,remaining)), allow_redirects=False, stream=True,
                          headers={'User-Agent': 'MiningInformationCollector/1.0'}) as response:
            if response.status_code in (301,302,303,307,308):
                url = urljoin(url, response.headers.get('Location', ''))
                continue
            if response.status_code in (401,403):
                raise ValueError('auth_required')
            response.raise_for_status()
            chunks, size = [], 0
            for chunk in response.iter_content(65536):
                if time.monotonic() > deadline:
                    raise ValueError('attachment_timeout')
                size += len(chunk)
                if size > MAX_BYTES:
                    raise ValueError('attachment_too_large')
                chunks.append(chunk)
            return b''.join(chunks), response.headers.get('Content-Type', '').split(';')[0]
    raise ValueError('redirect_limit')


def collect(response):
    output = []
    links = discover(response)
    deadline = time.monotonic() + 180
    for link in links[:4]:
        item = dict(link, status='failed', sha256='', mime_type='', byte_size=0, text='', pages=0,
                    fetched_at=datetime.now(timezone.utc).isoformat(timespec='seconds'))
        try:
            if time.monotonic() > deadline:
                raise ValueError('attachment_timeout')
            raw, mime = _download(link['url'], response.url, deadline=deadline)
            item.update(_raw=raw, sha256=hashlib.sha256(raw).hexdigest(), mime_type=mime, byte_size=len(raw))
            item.update(parse_content(raw, link['url'], timeout=deadline - time.monotonic()))
        except Exception as exc:
            code = str(exc) if isinstance(exc, ValueError) else type(exc).__name__
            status = 'unreadable' if isinstance(exc, WordConversionError) else 'failed'
            item.update(status='auth_required' if code == 'auth_required' else status, error=code[:80])
        output.append(item)
    if len(links) > 4:
        output.append(dict(name='其他附件',url=response.url,status='failed',sha256='',text='',
                           error='attachment_count_limit',fetched_at=datetime.now(timezone.utc).isoformat()))
    return output


def signature(items):
    return sorted((x['url'], x.get('sha256', ''), x['status']) for x in items)


def saved(db, record_id, version):
    row = db.execute('SELECT attachments_json FROM version_attachments WHERE record_id=? AND version=?',
                     (record_id,version)).fetchone()
    return json.loads(row['attachments_json']) if row else []


def save(db, root, record_id, version, items):
    for item in items:
        raw = item.pop('_raw', None)
        if raw is not None:
            extension = Path(urlsplit(item['url']).path).suffix.lower()
            if extension not in ('.pdf','.doc','.docx','.xlsx','.xls'):
                extension = '.bin'
            rel = Path('archive') / str(record_id) / f'v{version}-attachments' / (item['sha256'][:20] + extension)
            target = root / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(raw)
            item['archive_path'] = rel.as_posix()
    db.execute('INSERT OR REPLACE INTO version_attachments(record_id,version,attachments_json) VALUES (?,?,?)',
               (record_id,version,json.dumps(items,ensure_ascii=False)))
