import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from collector.attachments import discover, collect, parse_content
from collector.legacy_word import WordConversionError
from collector.db import connect, sync_sources
from collector.pipeline import ingest
from collector.report import export_jsonl

SOURCE = {'id':'public','name':'公开公告','category':'procurement','entry_url':'https://example.org',
          'adapter':'html_list','enabled':True,'access_status':'public','content_selector':'.body'}


class Response:
    url = 'https://example.org/notice.html'
    text = '''<div class="body">采购钢材；范围详见物料清单附件。</div>
    <div class="article-materials"><a onclick="ztbfjyz('/material.pdf','0','1')">物料清单附件.pdf</a>
    <a href="/plan.doc">采购方案.doc</a><a href="https://evil.org/x.pdf">外部</a></div>'''
    content = text.encode()
    headers = {'Content-Type':'text/html'}
    encoding = apparent_encoding = 'utf-8'


class AttachmentTests(unittest.TestCase):
    def test_outside_body_onclick_links_discovered(self):
        self.assertEqual(len(discover(Response())),2)
        self.assertEqual(discover(Response())[0]['url'],'https://example.org/material.pdf')

    def test_export_capture_and_attachment_only_change_preserve_versions(self):
        with tempfile.TemporaryDirectory() as folder, patch('collector.attachments.convert_legacy_word',
                                                           side_effect=WordConversionError('conversion_failed')):
            root = Path(folder)
            db = connect(root/'test.db')
            sync_sources(db,[SOURCE])
            run = db.execute("INSERT INTO runs(started_at,status) VALUES ('2026-10-10','running')").lastrowid
            entry = {'url':Response.url,'title':'钢材采购招标公告'}
            with patch('collector.attachments._download',return_value=(b'old doc', 'application/msword')):
                self.assertEqual(ingest(db,root,run,SOURCE,entry,Response()),'new')
                self.assertEqual(ingest(db,root,run,SOURCE,entry,Response()),'unchanged')
            first = db.execute('SELECT body_text FROM versions WHERE version=1').fetchone()[0]
            self.assertNotIn('采购方案.doc',first)
            with patch('collector.attachments._download',return_value=(b'changed doc','application/msword')):
                self.assertEqual(ingest(db,root,run,SOURCE,entry,Response()),'changed')
            self.assertEqual(db.execute('SELECT COUNT(*) FROM versions').fetchone()[0],2)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM version_attachments').fetchone()[0],2)
            export_jsonl(db,root/'out.jsonl')
            record = json.loads((root/'out.jsonl').read_text(encoding='utf-8'))
            self.assertEqual(len(record['attachments']),2)
            self.assertEqual(record['attachments'][0]['status'],'unsupported')
            self.assertTrue((root/record['attachments'][0]['archive_path']).exists())
            db.close()

    def test_doc_download_automatically_converts_and_archives_original(self):
        from types import SimpleNamespace
        import hashlib
        raw = b'original Word document'
        response = Response()
        response.text = '<a href="/plan.doc">plan.doc</a>'
        reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: '采购方案正文')])
        with patch('collector.attachments._download', return_value=(raw, 'application/msword')), \
                patch('collector.attachments.convert_legacy_word', return_value=b'%PDF-converted') as convert, \
                patch('collector.attachments.PdfReader', return_value=reader):
            result = collect(response)[0]
        convert.assert_called_once()
        self.assertEqual(result['status'], 'parsed')
        self.assertEqual(result['text'], '采购方案正文')
        self.assertEqual(result['pages'], 1)
        self.assertEqual(result['_raw'], raw)
        self.assertEqual(result['sha256'], hashlib.sha256(raw).hexdigest())

    def test_docx_paragraphs_and_tables_are_automatically_extracted(self):
        import io
        from docx import Document
        document = Document()
        document.add_paragraph('采购方案正文')
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = '钢材'
        output = io.BytesIO()
        document.save(output)
        result = parse_content(output.getvalue(), 'https://example.org/plan.docx')
        self.assertEqual(result['status'], 'parsed')
        self.assertIn('采购方案正文', result['text'])
        self.assertIn('钢材', result['text'])


if __name__ == '__main__': unittest.main()
